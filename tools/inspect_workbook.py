from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

workbook_path = Path(r"C:\Projects\life_os\_HABBIT TRACKER OF may 2026.xlsx")
wb = load_workbook(workbook_path, data_only=False)
wb_values = load_workbook(workbook_path, data_only=True)

print("sheets", wb.sheetnames)
for ws in wb.worksheets:
    ws_values = wb_values[ws.title]
    print("\nSHEET", ws.title, "dim", ws.max_row, ws.max_column)
    print("merged", [str(rng) for rng in list(ws.merged_cells.ranges)[:30]])
    charts = []
    for chart in ws._charts:
        title = None
        try:
            title = chart.title.tx.rich.p[0].r[0].t
        except Exception:
            title = str(getattr(chart, "title", None))
        charts.append((title, type(chart).__name__))
    print("charts", charts)

    for row_index in range(1, min(ws.max_row, 110) + 1):
        values = [ws.cell(row_index, col_index).value for col_index in range(1, min(ws.max_column, 24) + 1)]
        if any(value is not None for value in values):
            print(row_index, values)

    print("CALCULATED PREVIEW")
    for row_index in [10, 11, 12, 13, 27, 28, 29, 30, 31, 32, 33, 47, 51, 52, 53, 54, 55, 59, 60, 61, 62, 78, 79, 80, 81, 82]:
        values = [ws_values.cell(row_index, col_index).value for col_index in range(1, min(ws.max_column, 42) + 1)]
        if any(value is not None for value in values):
            print("V", row_index, values)

    formulas = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append((cell.coordinate, cell.value))
    print("formula_count", len(formulas))
    for coordinate, formula in formulas[:160]:
        print("F", coordinate, formula)

    fills = []
    fonts = []
    dimensions = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 35), min_col=1, max_col=min(ws.max_column, 24)):
        for cell in row:
            if cell.value is not None:
                fills.append((cell.coordinate, cell.fill.fill_type, cell.fill.fgColor.type, cell.fill.fgColor.rgb, cell.fill.fgColor.indexed))
                font_color = cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else None
                fonts.append((cell.coordinate, cell.font.name, cell.font.sz, cell.font.bold, font_color))
    for idx in range(1, min(ws.max_column, 50) + 1):
        letter = get_column_letter(idx)
        dimensions.append((letter, ws.column_dimensions[letter].width))
    print("fills sample", fills[:100])
    print("fonts sample", fonts[:60])
    print("column widths", dimensions)

    print("row heights", [(idx, ws.row_dimensions[idx].height) for idx in range(1, min(ws.max_row, 110) + 1) if ws.row_dimensions[idx].height])
